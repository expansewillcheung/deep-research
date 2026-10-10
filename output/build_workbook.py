# -*- coding: utf-8 -*-
"""Build IB-style US sector valuation & stock selection workbook (data as of 2026-10-09 close)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule

OUT = "/home/user/deep-research/output/US_Sector_Valuation_Stock_Selection_2026-10.xlsx"

# ---------------- styles ----------------
NAVY = "1F3864"; MID = "2F5597"; LIGHT = "D9E1F2"; GREY = "F2F2F2"; GOLD = "FFF2CC"
F = "Arial"
def font(color="000000", bold=False, size=10, italic=False):
    return Font(name=F, color=color, bold=bold, size=size, italic=italic)
BLUE = font("0000FF"); BLACK = font(); GREEN = font("008000")
HDR = font("FFFFFF", True); TITLE = font(NAVY, True, 16); SUB = font(MID, True, 11)
NOTE = font("595959", False, 8, True)
fill = lambda c: PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM = Border(bottom=Side(style="medium", color=NAVY))
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

PCT = '0.0%;(0.0%);"-"'
MULT = '0.0"x";(0.0"x");"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
USD0 = '$#,##0;($#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
INT = '0'

wb = Workbook()

def sheet(name, title, subtitle, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws["A1"] = title; ws["A1"].font = TITLE
    ws["A2"] = subtitle; ws["A2"].font = NOTE
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
    return ws

def header(ws, row, labels, col=1, height=32):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = HDR; c.fill = fill(NAVY); c.alignment = CENTER; c.border = BOX
    ws.row_dimensions[row].height = height

def section(ws, row, text, ncol):
    c = ws.cell(row=row, column=1, value=text); c.font = font("FFFFFF", True, 11)
    for i in range(1, ncol + 1):
        ws.cell(row=row, column=i).fill = fill(MID)

def put(ws, row, col, val, f=BLACK, fmt=None, al=None, bd=True, fl=None):
    c = ws.cell(row=row, column=col, value=val)
    c.font = f
    if fmt: c.number_format = fmt
    c.alignment = al if al else Alignment(vertical="center")
    if bd: c.border = BOX
    if fl: c.fill = fill(fl)
    return c

def band(ws, r0, r1, c0, c1):
    for r in range(r0, r1 + 1):
        if (r - r0) % 2 == 1:
            for c in range(c0, c1 + 1):
                ws.cell(row=r, column=c).fill = fill(GREY)

# ======================= DATA =======================
TICK = ["BRK.B", "JPM", "V", "NEE", "CEG", "SO", "WELL", "PLD", "EQIX", "LIN", "NEM", "FCX"]
NAME = {"BRK.B": "Berkshire Hathaway", "JPM": "JPMorgan Chase", "V": "Visa", "NEE": "NextEra Energy",
        "CEG": "Constellation Energy", "SO": "Southern Co.", "WELL": "Welltower", "PLD": "Prologis",
        "EQIX": "Equinix", "LIN": "Linde", "NEM": "Newmont", "FCX": "Freeport-McMoRan"}
SECT = {"BRK.B": "金融", "JPM": "金融", "V": "金融", "NEE": "公用事業", "CEG": "公用事業", "SO": "公用事業",
        "WELL": "房地產", "PLD": "房地產", "EQIX": "房地產", "LIN": "原材料", "NEM": "原材料", "FCX": "原材料"}
SUBI = {"BRK.B": "多元控股/保險", "JPM": "多元化銀行", "V": "支付網絡", "NEE": "受監管電力+再生能源",
        "CEG": "獨立發電(核電)", "SO": "受監管電力", "WELL": "醫療/長者住宅REIT", "PLD": "物流REIT",
        "EQIX": "數據中心REIT", "LIN": "工業氣體", "NEM": "黃金開採", "FCX": "銅金開採"}
# price, mcap $B, PE ttm, PB, DY%, YTD%
MKT = {"BRK.B": (515.62, 1112.126, 12.97, 1.49, 0.00, 2.58), "JPM": (332.99, 885.149, 13.91, 2.50, 1.85, 5.33),
       "V": (385.45, 741.555, 33.11, 21.39, 0.70, 10.57), "NEE": (77.38, 161.413, 17.36, 2.83, 3.15, -1.51),
       "CEG": (298.07, 105.608, 30.48, 3.30, 0.56, -15.25), "SO": (86.17, 99.127, 21.28, 2.51, 3.48, 1.25),
       "WELL": (225.47, 162.506, 104.80, 3.50, 1.36, 22.78), "PLD": (129.49, 120.715, 28.72, 2.25, 3.26, 3.81),
       "EQIX": (1026.06, 101.243, 66.04, 7.04, 1.97, 35.97), "LIN": (483.73, 222.990, 30.78, 5.71, 1.30, 14.57),
       "NEM": (117.84, 124.167, 14.44, 3.52, 0.87, 18.78), "FCX": (73.69, 105.820, 36.04, 5.26, 0.51, 46.12)}
FWD = {"BRK.B": 22.163, "JPM": 24.46, "V": 14.493, "NEE": 4.155, "CEG": 12.041, "SO": 4.633,
       "WELL": 2.83, "PLD": 3.48, "EQIX": 17.015, "LIN": 18.74, "NEM": 8.99, "FCX": 3.80}
# valuation band: metric, now, low, median, high, pct-time-cheaper, window, industry median
BANDV = {"BRK.B": ("P/B", 1.49, 1.38, 1.46, 1.59, .441, "5Y", 1.48), "JPM": ("P/E", 13.91, 11.37, 13.07, 14.84, .4043, "3Y", 11.99),
         "V": ("P/E", 33.11, 29.01, 31.33, 33.13, .2015, "3Y", 14.11), "NEE": ("P/E", 17.36, 19.15, 25.22, 42.59, .8666, "5Y", 18.73),
         "CEG": ("P/E", 30.48, 23.77, 28.93, 37.11, .4707, "3Y", 18.73), "SO": ("P/E", 21.28, 20.52, 22.49, 25.05, .6894, "5Y", 18.73),
         "WELL": ("P/B", 3.50, 1.82, 2.19, 3.20, .0782, "5Y", 1.36), "PLD": ("P/B", 2.25, 1.94, 2.11, 2.45, .3574, "5Y", 1.36),
         "EQIX": ("P/B", 7.04, 5.53, 5.99, 6.70, .0867, "5Y", 1.36), "LIN": ("P/E", 30.78, 32.29, 33.68, 39.91, .908, "5Y", 29.45),
         "NEM": ("P/E", 14.44, 13.10, 14.92, 17.30, .5632, "1Y", 8.53), "FCX": ("P/E", 36.04, 29.07, 32.55, 37.75, .2679, "3Y", 31.20)}
# analyst: target mean, hi, lo, buy, over, hold, under, sell, total
RAT = {"BRK.B": (547.67, 604, 510, 2, 0, 2, 0, 0, 4), "JPM": (372.68, 420, 305, 9, 4, 12, 0, 0, 26),
       "V": (419.63, 466, 330, 27, 9, 3, 1, 0, 41), "NEE": (97.05, 114, 56, 11, 3, 6, 0, 1, 22),
       "CEG": (341.53, 395, 290, 13, 6, 3, 0, 0, 22), "SO": (98.13, 112, 79, 6, 1, 14, 0, 2, 23),
       "WELL": (262.05, 292, 219, 12, 6, 4, 0, 0, 22), "PLD": (158.23, 170, 140, 7, 3, 7, 0, 0, 18),
       "EQIX": (1234.19, 1380, 1060, 23, 5, 6, 0, 0, 34), "LIN": (543.04, 580, 400, 17, 5, 5, 0, 1, 28),
       "NEM": (137.33, 170, 67, 15, 5, 2, 0, 1, 23), "FCX": (77.55, 92, 32, 12, 6, 4, 0, 1, 24)}
# financials FY2025..FY2020 : rev $B, NI $B, EPS, ROE %, OCF $B, CapEx $B (None if n/a)
FIN = {
 "BRK.B": ([371.444, 371.433, 364.482, 302.020, 276.185, 245.579], [66.968, 88.995, 96.223, -22.759, 89.937, 42.521],
           [31.04, 41.27, 44.27, -10.33, 39.70, 17.78], [9.80, 14.70, 18.60, -4.65, 18.95, 9.80],
           [45.969, 30.592, 49.196, 37.350, 39.427, 39.773], [20.927, 18.976, 19.409, 15.464, 13.276, 13.012]),
 "JPM": ([168.235, 166.775, 145.670, 122.306, 130.898, 102.471], [55.681, 56.868, 47.760, 35.892, 46.503, 27.410],
         [20.02, 19.75, 16.23, 12.09, 15.36, 8.88], [16.69, 18.19, 16.89, 13.69, 18.29, 11.34], None, None),
 "V": ([40.000, 35.926, 32.653, 29.310, 24.105, 21.846], [19.853, 19.457, 16.989, 14.630, 12.008, 10.475],
       [10.20, 9.73, 8.28, 7.00, 5.63, 4.89], [52.75, 51.79, 48.34, 43.18, 36.59, 34.72],
       [23.059, 19.950, 20.755, 18.849, 15.227, 10.440], [1.482, 1.257, 1.059, 0.970, 0.705, 0.736]),
 "NEE": ([27.412, 24.753, 28.114, 20.956, 17.069, 17.997], [6.835, 6.946, 7.310, 4.147, 3.573, 2.919],
         [3.30, 3.37, 3.60, 2.10, 1.81, 1.48], [13.06, 14.24, 16.86, 10.85, 9.69, 7.94],
         [12.485, 13.260, 11.301, 8.262, 7.553, 7.983], None),
 "CEG": ([25.533, 23.568, 24.918, 24.440, 19.649, 17.603], [2.319, 3.749, 1.623, -0.160, -0.205, 0.589],
         [7.40, 11.89, 5.01, -0.49, None, None], [16.75, 31.12, 14.79, -1.44, -1.74, 4.55],
         [4.237, -2.464, -5.301, -2.353, -1.338, 0.584], None),
 "SO": ([29.553, 26.724, 25.253, 29.279, 23.113, 20.375], [4.341, 4.401, 3.976, 3.524, 2.393, 3.119],
        [3.92, 3.99, 3.62, 3.26, 2.24, 2.93], [12.54, 13.61, 12.86, 12.09, 8.57, 11.24],
        [9.802, 9.788, 7.553, 6.302, 6.169, 6.696], None),
 "WELL": ([10.838, 7.991, 6.573, 5.861, 4.737, 4.587], [0.937, 0.952, 0.340, 0.141, 0.336, 0.979],
          [1.39, 1.57, 0.66, 0.30, 0.78, 2.33], [2.53, 3.32, 1.49, 0.74, 2.00, 6.21],
          [2.882, 2.256, 1.602, 1.329, 1.275, 1.365], None),
 "PLD": ([9.193, 8.555, 8.331, 6.285, 5.164, 4.736], [3.322, 3.726, 3.053, 3.359, 2.934, 1.473],
         [3.56, 4.01, 3.29, 4.25, 3.94, 2.01], [6.21, 6.96, 5.75, 7.76, 8.99, 5.41],
         [5.008, 4.912, 5.373, 4.126, 2.996, 2.937], None),
 "EQIX": ([9.260, 8.737, 8.197, 7.117, 6.260, 5.623], [1.350, 0.815, 0.969, 0.705, 0.500, 0.370],
          [13.76, 8.50, 10.31, 7.67, 5.53, 4.18], [9.75, 6.27, 8.08, 6.30, 4.65, 3.80],
          [3.911, 3.249, 3.217, 2.963, 2.547, 2.310], None),
 "LIN": ([33.986, 33.005, 32.854, 33.364, 30.793, 27.243], [6.898, 6.565, 6.199, 4.147, 3.826, 2.501],
         [14.61, 13.62, 12.59, 8.23, 7.33, 4.71], [18.07, 16.87, 15.55, 9.87, 8.37, 5.18],
         [10.350, 9.423, 9.305, 8.864, 9.725, 7.429], [5.261, 4.497, 3.787, 3.173, 3.086, 3.400]),
 "NEM": ([22.669, 18.682, 11.812, 11.915, 12.222, 11.497], [7.085, 3.348, -2.494, -0.429, 1.166, 2.829],
         [6.39, 2.92, -2.97, -0.54, 1.46, 3.51], [22.21, 11.13, -10.42, -2.22, 4.93, 12.00],
         [10.334, 6.363, 2.763, 3.220, 4.279, 4.882], [3.035, 3.402, 2.666, 2.131, 1.653, 1.302]),
 "FCX": ([25.915, 25.455, 22.855, 22.780, 22.845, 14.198], [2.197, 1.883, 1.842, 3.461, 4.299, 0.596],
         [1.52, 1.30, 1.28, 2.39, 2.90, 0.41], [12.04, 10.99, 11.42, 23.44, 35.60, 6.12],
         [5.610, 7.160, 5.279, 5.139, 7.715, 3.017], [4.494, 4.808, 4.824, 3.469, 2.115, 1.961]),
}
BS = {"BRK.B": (1222.176, 502.473), "JPM": (4424.900, 4062.462), "V": (99.627, 61.718), "NEE": (212.721, 146.242),
      "CEG": (57.249, 42.396), "SO": (155.720, 116.853), "WELL": (67.303, 24.100), "PLD": (98.724, 40.970),
      "EQIX": (40.141, 25.963), "LIN": (86.817, 47.076), "NEM": (57.121, 23.079), "FCX": (58.167, 27.401)}

# ======================= COVER =======================
ws0 = wb.active; ws0.title = "Cover"; ws0.sheet_view.showGridLines = False
ws0.column_dimensions["A"].width = 3; ws0.column_dimensions["B"].width = 110
for r in range(1, 40):
    for c in range(1, 4):
        ws0.cell(row=r, column=c).fill = fill(NAVY)
ws0["B4"] = "美股板塊估值地圖 & 低估板塊龍頭選股分析"; ws0["B4"].font = font("FFFFFF", True, 24)
ws0["B6"] = "US Sector Valuation Map & Value-Sector Leaders: 5-Year Buffett-Lens Stock Selection"; ws0["B6"].font = font("BDD7EE", False, 14)
ws0["B9"] = "數據截至：2026-10-09 美股收盤（LongBridge OpenAPI）｜板塊估值：FactSet Earnings Insight（2026-10-09）"; ws0["B9"].font = font("FFFFFF", False, 11)
ws0["B10"] = "覆蓋：GICS 11 大板塊 → 4 個低估/低估值板塊 → 12 家龍頭（每板塊市值前 3）→ 最終 5+1 檔 5 年期組合"; ws0["B10"].font = font("FFFFFF", False, 11)
toc = [("Summary", "結論摘要與最終組合"), ("Sector_Map", "11 大板塊 Forward P/E vs 5Y/10Y 均值、盈利增速、低估/高估判定"),
       ("Methodology", "各板塊估值錨、P/E 失真點、板塊估值算法合理性評估"), ("Competition", "板塊內部競爭方式、集中度、護城河來源"),
       ("Macro", "宏觀儀表板 + 個股宏觀敏感度矩陣"), ("Comps", "12 家龍頭可比估值表（可重算）"),
       ("Financials", "FY2020–FY2025 財務數據（LongBridge）"), ("Valuation", "多方法估值：歷史倍數 / 同業倍數 / 兩階段 DCF·DDM / 共識目標價 + 反向 DCF"),
       ("Scorecard", "巴菲特準則：業務易懂、護城河、ROE、資產負債表、管理層、宏觀韌性 + 估值分"),
       ("Picks", "5 年期預期回報、排序、組合權重、買入理由、出場條件"), ("Sources", "數據來源、假設、免責聲明")]
ws0["B13"] = "目錄 / Contents"; ws0["B13"].font = font("FFD966", True, 13)
for i, (s, d) in enumerate(toc):
    c = ws0.cell(row=14 + i, column=2, value=f"{i+1:>2}.  {s}  —  {d}")
    c.font = font("FFFFFF", False, 11); c.hyperlink = f"#{s}!A1"
ws0["B27"] = "圖例：藍字 = 硬編碼輸入/數據；黑字 = 公式；綠字 = 跨表連結；黃底 = 關鍵假設（可改動，全表自動重算）"; ws0["B27"].font = font("FFD966", False, 10)
ws0["B29"] = "免責聲明：本工作簿僅供研究與教育用途，不構成任何投資建議或買賣邀約。所有估值依賴公開數據與主觀假設，可能有誤；投資前請自行核實並諮詢持牌顧問。"
ws0["B29"].font = font("BFBFBF", False, 9, True); ws0["B29"].alignment = WRAP

# ======================= SECTOR MAP =======================
S = sheet("Sector_Map", "板塊估值地圖 — S&P 500 11 大板塊", "來源：FactSet Earnings Insight 2026-10-09（p.29/30/34/37）；10Y 殖利率連結 Macro!C8。藍字為原始數據，黑字為公式。",
          [20, 8, 9, 9, 9, 10, 10, 10, 9, 9, 10, 10, 9, 9, 10, 12, 13, 46, 14])
cols = ["板塊", "ETF", "Fwd P/E", "5Y 均值", "10Y 均值", "vs 10Y", "vs 5Y", "CY26 EPS 增速", "CY27 EPS 增速", "CY27 營收增速",
        "PEG (CY27)", "盈利收益率", "EY − 10Y", "目標價上行", "Buy 評級%", "規則判定 (vs10Y ±5%)", "分析師修正判定", "關鍵理由 / 估值走向", "深挖?"]
header(S, 4, cols, height=40)
SEC = [("可選消費 Cons. Disc.", "XLY", 23.2, 26.0, 25.5, .352, -.025, .084, .240, .58, "合理偏貴", "歷史均值被 AMZN/TSLA 高倍數抬高；CY27 EPS −2.5%，估值下修中→「低於均值」不等於便宜", "否"),
       ("工業 Industrials", "XLI", 22.8, 20.9, 20.0, .121, .163, .083, .247, .58, "高估", "國防/電網設備/AI 資本開支題材溢價，較 10Y 均值 +14%；趨勢：溢價擴張", "否"),
       ("必需消費 Staples", "XLP", 21.6, 20.8, 20.1, .075, .069, .042, .119, .47, "高估", "CY27 EPS 僅 +6.9% 卻 21.6x，PEG 最差；盈利遭下修（WMT/TSN）", "否"),
       ("資訊科技 Info. Tech.", "XLK", 21.3, 25.6, 23.1, .534, .416, .292, .214, .69, "合理（週期頂部風險）", "低 P/E 源於半導體盈利 +130%；PEG 僅 0.5，但若 AI 資本開支放緩 EPS 下修風險大", "否"),
       ("S&P 500", "SPY", 19.1, 19.8, 19.1, .327, .158, .094, .196, .60, "合理", "Fwd 19.1x = 10Y 均值；Trailing 26.0x 偏高；盈利動能強（Q3 +29.6%）", "—"),
       ("通訊服務 Comm. Svcs", "XLC", 18.6, 18.3, 17.6, .582, -.112, .146, .205, .72, "合理偏貴", "CY27 EPS −11.2%（Meta 高基數），表面 P/E 低估盈利回落", "否"),
       ("醫療保健 Health Care", "XLV", 18.3, 17.4, 16.6, .019, .225, .043, .139, .61, "結構分化", "LLY(41x) 拉高板塊；UNH/MRK/PFE 等權看偏便宜；CY27 EPS 反彈 +22.5%", "觀察"),
       ("房地產 Real Estate", "XLRE", 16.7, 18.1, 18.6, .101, .064, .079, .235, .55, "低估（需用 P/FFO 驗證）", "加息壓制（Fed 9 月加至 3.75–4.00%）；P/E 對 REIT 失真，PLD 合理、WELL 偏貴", "是"),
       ("原材料 Materials", "XLB", 16.6, 17.9, 17.8, .380, .121, .060, .222, .63, "低估（週期）", "金/銅價高位推升盈利；Q3 盈利預期 −9.6% 下修；龍頭 LIN 處 5 年估值低位", "是"),
       ("公用事業 Utilities", "XLU", 16.1, 18.0, 18.0, .130, .089, .049, .211, .53, "低估", "利率上行令股息吸引力下降（NEE 半年 −17%）；AI 數據中心用電需求是長期支撐", "是"),
       ("金融 Financials", "XLF", 14.3, 14.6, 13.9, .147, .089, .053, .158, .55, "合理（絕對值低）", "絕對 P/E 第二低；銀行應看 P/B-ROE；加息利好 NII，信貸週期是風險", "是"),
       ("能源 Energy", "XLE", 13.5, 12.6, 11.0, .966, -.093, -.052, .088, .64, "高估（盈利頂部）", "低 P/E 是週期陷阱：油價 $90（伊朗衝突溢價）推高盈利，CY27 EPS −9.3%，較 10Y 均值 +23%", "否")]
r0 = 5
for i, s in enumerate(SEC):
    r = r0 + i
    put(S, r, 1, s[0], font(bold=(s[0] == "S&P 500")))
    put(S, r, 2, s[1], BLUE, al=CENTER)
    for j, v in enumerate(s[2:5]):
        put(S, r, 3 + j, v, BLUE, MULT)
    put(S, r, 6, f"=C{r}/E{r}-1", fmt=PCT)
    put(S, r, 7, f"=C{r}/D{r}-1", fmt=PCT)
    put(S, r, 8, s[5], BLUE, PCT); put(S, r, 9, s[6], BLUE, PCT); put(S, r, 10, s[7], BLUE, PCT)
    put(S, r, 11, f'=IF(I{r}>0,C{r}/(I{r}*100),"n/m")', fmt=NUM2)
    put(S, r, 12, f"=1/C{r}", fmt=PCT)
    put(S, r, 13, f"=L{r}-Macro!$C$8", fmt=PCT)
    put(S, r, 14, s[8], BLUE, PCT); put(S, r, 15, s[9], BLUE, PCT)
    put(S, r, 16, f'=IF(F{r}<=-0.05,"低估",IF(F{r}>=0.05,"高估","合理"))', al=CENTER)
    put(S, r, 17, s[10], font(bold=True), al=CENTER)
    put(S, r, 18, s[11], al=WRAP)
    put(S, r, 19, s[12], al=CENTER, fl=(GOLD if s[12] == "是" else None))
    S.row_dimensions[r].height = 42
band(S, r0, r0 + len(SEC) - 1, 1, 15)
S.conditional_formatting.add(f"F{r0}:G{r0+11}", ColorScaleRule(start_type="num", start_value=-0.15, start_color="63BE7B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="num", end_value=0.25, end_color="F8696B"))
S.conditional_formatting.add(f"P{r0}:P{r0+11}", CellIsRule(operator="equal", formula=['"低估"'], fill=fill("C6EFCE"), font=font("006100", True)))
S.conditional_formatting.add(f"P{r0}:P{r0+11}", CellIsRule(operator="equal", formula=['"高估"'], fill=fill("FFC7CE"), font=font("9C0006", True)))
r = r0 + len(SEC) + 1
section(S, r, "板塊排序（按 vs 10Y 折讓，公式自動）", 19)
header(S, r + 1, ["排名", "板塊", "vs 10Y", "Fwd P/E", "PEG"], height=22)
for k in range(1, 13):
    rr = r + 1 + k
    put(S, rr, 1, k, al=CENTER)
    put(S, rr, 2, f"=INDEX($A${r0}:$A${r0+11},MATCH(SMALL($F${r0}:$F${r0+11},A{rr}),$F${r0}:$F${r0+11},0))")
    put(S, rr, 3, f"=SMALL($F${r0}:$F${r0+11},A{rr})", fmt=PCT)
    put(S, rr, 4, f"=INDEX($C${r0}:$C${r0+11},MATCH(C{rr},$F${r0}:$F${r0+11},0))", fmt=MULT)
    put(S, rr, 5, f"=INDEX($K${r0}:$K${r0+11},MATCH(C{rr},$F${r0}:$F${r0+11},0))", fmt=NUM2)
rr = r + 15
section(S, rr, "選擇深挖板塊的邏輯", 19)
txt = ["1) 以「Fwd P/E 相對 10Y 均值」作首輪篩選：公用事業(−10.6%)、房地產(−10.2%)、可選消費(−9.0%)、資訊科技(−7.8%)、原材料(−6.7%) 為表面低估。",
       "2) 剔除「假便宜」：可選消費與資訊科技的折讓來自盈利暴增/集中度，不是價格便宜（盈利頂部風險）；能源絕對 P/E 最低卻高於歷史 23%，是典型週期陷阱。",
       "3) 加入金融：絕對 P/E 14.3x 為第二低、略高於 10Y 均值（+2.9%），但 P/B-ROE 框架下龍頭仍合理，且加息環境利好淨息差，屬「低估值板塊」。",
       "4) 最終深挖：金融、公用事業、房地產、原材料 4 板塊 × 市值前 3 = 12 家龍頭（市值取自 LongBridge 2026-10-09）。"]
for i, t in enumerate(txt):
    c = S.cell(row=rr + 1 + i, column=1, value=t); c.font = BLACK
    S.merge_cells(start_row=rr + 1 + i, start_column=1, end_row=rr + 1 + i, end_column=18)
S.freeze_panes = "B5"

# ======================= METHODOLOGY =======================
M = sheet("Methodology", "估值方法論 & 板塊估值算法合理性評估", "各板塊應使用的估值錨，以及「單看 P/E」在何處失真。P/E 適用度：5=很適用，1=基本不適用（分析師判斷）。",
          [20, 30, 30, 44, 50, 10])
header(M, 4, ["板塊/子行業", "主估值錨", "輔助方法", "為什麼（經濟邏輯）", "單看 P/E 的失真點", "P/E 適用度"], height=30)
MET = [("銀行 (JPM)", "P/B vs ROE：合理 P/B=(ROE−g)/(r−g)；P/TBV", "P/E、股息+回購收益率、CET1 超額資本", "銀行價值來自權益資本的超額回報；資本受監管約束", "撥備計提/釋放令 EPS 週期性大起大落；2022 年 P/E 失真", 3),
       ("保險/綜合控股 (BRK)", "P/B、內在價值(營運盈利×倍數+超額現金)", "浮存金成本、分部加總(SOTP)", "巴菲特本人以 P/B 及回購門檻衡量；GAAP 淨利含未實現投資損益", "GAAP EPS 2022 年 −$10.3、2023 年 $44.3，P/E 毫無意義", 1),
       ("支付網絡 (V/MA)", "P/E、FCF 收益率、PEG", "DCF（高轉換率 FCF≈淨利）", "輕資產、FCF 轉換率 >100%、增長可預測", "訴訟撥備(每季 $0.6–1.0B) 壓低 GAAP EPS；需用調整後", 5),
       ("受監管公用事業 (NEE/SO)", "Fwd P/E vs 10Y 殖利率；股息折現(DDM)", "費率基礎(rate base)增長 × 允許 ROE", "盈利=允許回報×費率基礎；類債券資產，對利率敏感", "一次性項目(NEE GAAP vs 調整後差異大)；忽略負 FCF 與股權攤薄", 4),
       ("獨立發電 (CEG)", "EV/EBITDA、FCF 收益率", "P/E、合約(PPA)覆蓋率", "收入隨電價/產能拍賣波動；數據中心 PPA 決定估值上限", "公允價值調整與 Calpine 併購攤銷令 GAAP 失真", 3),
       ("REITs (WELL/PLD/EQIX)", "P/FFO、P/AFFO、NAV/隱含資本化率", "股息收益率 vs 10Y、DCF", "房地產折舊非真實經濟成本；REIT 須派 90% 應稅收入", "WELL GAAP P/E 105x vs P/FFO 35x；EQIX 66x vs P/AFFO 24x — P/E 完全失真", 1),
       ("工業氣體 (LIN)", "P/E、EV/EBITDA", "ROC、合約 backlog", "15–20 年照付不議合約，盈利穩定如公用事業但增長更快", "較少失真；需看調整後 EPS（剔除俄羅斯減值等）", 5),
       ("金屬礦業 (NEM/FCX)", "P/NAV、中週期價格下 EV/EBITDA", "在期貨曲線下的 FCF 收益率、AISC 曲線位置", "價格接受者，價值=儲量×(價格−成本)", "金價 $4,500 / 銅價 $6.5 時 P/E 最低 → 週期頂部陷阱", 2),
       ("能源 (XOM/CVX)", "中週期油價($65–70)下 EV/EBITDA、FCF 收益率", "P/CF、儲量替代率", "盈利與油價高度相關", "油價 $90 時 Fwd P/E 13.5x 看似便宜，CY27 EPS −9.3%", 2),
       ("資訊科技", "成熟：P/E；高增長：EV/Sales、PEG", "Rule of 40、DCF", "增長率差異極大；半導體有強週期", "半導體盈利 +130% 令板塊 P/E 低於均值，週期頂部風險", 3),
       ("醫療保健", "P/E + 管線 rNPV", "專利懸崖調整後 EPS", "專利到期導致盈利斷崖", "LLY 一家拉高板塊；收購研發費用一次性扣除(MRK P/E 113x)", 3),
       ("必需消費", "P/E、股息率 vs 債券", "EV/EBITDA", "低增長、類債券", "低增長配高倍數時 P/E 看似「合理」但 PEG 很差", 4)]
for i, m in enumerate(MET):
    r = 5 + i
    for j, v in enumerate(m):
        put(M, r, 1 + j, v, BLUE if j == 5 else BLACK, al=(CENTER if j == 5 else WRAP))
    M.row_dimensions[r].height = 46
band(M, 5, 5 + len(MET) - 1, 1, 6)
r = 5 + len(MET) + 1
section(M, r, "板塊估值算法（FactSet 加總 Fwd P/E）合理性評估", 6)
CRIT = [("市值加權偏差", "板塊 P/E = 板塊總市值 ÷ 總預期盈利，被巨頭主導（IT 前 3 大佔比過半；醫療被 LLY 主導）。", "高", "同時看等權 P/E 與中位數（本表 Comps 用個股中位數交叉驗證）"),
        ("一致預期樂觀偏差", "歷史上分析師季內平均下修 EPS 2–4%；本季反而上修 1.4%，盈利動能強但頂部特徵明顯。", "中", "對週期股用中週期/正常化盈利"),
        ("週期頂部陷阱", "能源/原材料/半導體在盈利高峰時 P/E 最低，按 P/E 會買在頂部。", "高", "用 EV/EBITDA@中週期價格、P/NAV"),
        ("利率環境未調整", "10Y 均值含 2020–21 零利率時期；現 10Y ≈4.6%、Fed 加息至 3.75–4.00%，合理倍數中樞應低於歷史均值。", "高", "看盈利收益率−10Y 的股權風險溢價（Sector_Map M 欄）"),
        ("GICS 分類變更", "2018 年 GOOGL/META 移入通訊；2023 年 V/MA 移入金融 → 金融板塊 P/E 歷史不可比（結構性上移）。", "中", "金融板塊 P/E 略高於 10Y 均值部分源於此，不代表變貴"),
        ("會計口徑", "REIT 折舊、保險未實現損益、併購攤銷令 GAAP P/E 失真。", "高", "各行業採用專屬指標（見上表）"),
        ("結論", "FactSet 板塊 Fwd P/E 適合作「首輪篩選」，不足以作買入依據；本工作簿以個股層面多方法估值 + 歷史分位 + 反向 DCF 交叉驗證。", "—", "—")]
header(M, r + 1, ["偏差來源", "說明", "影響", "本分析的修正方式"], height=22)
for i, c_ in enumerate(CRIT):
    rr = r + 2 + i
    put(M, rr, 1, c_[0], font(bold=True), al=WRAP); put(M, rr, 2, c_[1], al=WRAP); put(M, rr, 3, c_[2], al=CENTER); put(M, rr, 4, c_[3], al=WRAP)
    M.merge_cells(start_row=rr, start_column=4, end_row=rr, end_column=5)
    M.row_dimensions[rr].height = 44
M.column_dimensions["B"].width = 46
M.freeze_panes = "B5"

# ======================= COMPETITION =======================
C = sheet("Competition", "板塊內部競爭方式 & 護城河來源", "分析師判斷 + 公司公開資料；市值為 LongBridge 2026-10-09。",
          [16, 40, 30, 26, 46])
header(C, 4, ["板塊/子行業", "競爭方式（靠什麼贏）", "集中度/格局", "市值前 3（$B）", "護城河來源 & 龍頭優勢"], height=30)
COMP = [("金融—銀行", "存款成本(低成本存款基礎)、規模攤薄科技投入(JPM 年科技支出~$18B)、監管資本作門檻、交叉銷售", "寡頭化：前 4 大銀行佔美國存款約 4 成；大而不倒隱性擔保", "BRK 1,112 / JPM 885 / V 742", "JPM：全業務線前 3、資產負債表堡壘、ROE 16–18%；V：雙邊網絡效應(持卡人×商戶)，與 MA 形成雙寡頭；BRK：低成本浮存金 + 永久資本 + 收購對象首選買家"),
        ("金融—支付", "網絡效應、受理覆蓋、品牌信任、安全/風控；新威脅：穩定幣、A2A 即時支付、CCCA 法案", "V/MA 雙寡頭（Amex/Discover 次之）", "V 742 / MA 516", "邊際成本近零、毛利 97.8%、ROE >50%"),
        ("公用事業—受監管", "不爭客戶，爭「費率基礎增長 + 監管關係 + 資本成本」；大型負荷(數據中心)合約成新戰場", "地區性壟斷", "NEE 161 / CEG 106 / SO 99", "NEE：FPL 最優營運效率 + 全美最大再生能源開發商；SO：東南部 17GW 大型負荷合約（含 OpenAI 3.2GW）；CEG：全美最大核電機組群(稀缺零碳基載)"),
        ("公用事業—獨立發電", "發電成本曲線位置、核電/燃氣稀缺性、與超大規模雲廠商的長期 PPA", "分散，但核電高度集中於 CEG/VST", "CEG 106 / VST 54", "核電執照與既有站點不可複製"),
        ("房地產—REITs", "地段 + 規模 + 資本成本 + 營運平台；數據中心看電力與互聯生態", "子行業內寡頭（物流 PLD、數據中心 EQIX/DLR、醫療 WELL/VTR）", "WELL 163 / PLD 121 / EQIX 101", "PLD：全球物流 ~12 億平方呎、客戶網絡、數據中心電力管線 5.8GW；EQIX：互聯網絡效應(客戶間直連)；WELL：長者住宅數據化營運平台、SHO 同店 NOI +20.5%"),
        ("原材料—工業氣體", "現場製氣長約(15–20 年照付不議)、管道網絡區域密度、轉換成本高", "三寡頭：Linde / Air Liquide / Air Products", "LIN 223", "LIN：區域密度 + 合約 backlog $11B + 電子特氣；定價權可傳導通脹"),
        ("原材料—金屬礦業", "成本曲線位置(AISC 分位)、儲量壽命、司法管轄區風險；無定價權", "分散；價格接受者", "NEM 124 / FCX 106", "NEM：全球最大金礦商、AISC $1,621/oz；FCX：Grasberg/美洲一級銅礦，但 2025 泥石流事故後產能 65%→2027 年底恢復"),
        ("對照—能源", "上游成本、儲量、資本紀律；下游煉油價差", "XOM/CVX 寡頭", "XOM 695 / CVX 419", "週期性強；本輪不深挖（盈利頂部）")]
for i, cc in enumerate(COMP):
    r = 5 + i
    for j, v in enumerate(cc):
        put(C, r, 1 + j, v, al=WRAP)
    C.row_dimensions[r].height = 62
band(C, 5, 5 + len(COMP) - 1, 1, 5)
C.freeze_panes = "B5"

# ======================= MACRO =======================
Mc = sheet("Macro", "宏觀儀表板 & 個股宏觀敏感度", "來源：LongBridge macrodata（CPI、失業率）；CNBC/Fed（政策利率）；FactSet（油價、盈利）；Webull/FCX（金價）。黃底=可調關鍵假設。",
           [30, 16, 14, 14, 14, 14, 14, 14, 12, 40])
section(Mc, 4, "A. 關鍵宏觀指標", 10)
header(Mc, 5, ["指標", "最新期", "數值", "前值", "趨勢", "來源"], height=22)
MAC = [("CPI 同比 (%)", "2026-08", .034, .034, "高位橫行（5 月峰值 4.2%）", "LongBridge 30771871"),
       ("核心 CPI 同比 (%)", "2026-05", .029, None, "溫和", "BLS via 媒體"),
       ("10Y 美債殖利率 (估)", "2026-09", .046, None, "4.6–4.8% 區間（關鍵假設）", "Fed H.15 / YCharts (7/24 4.66%, 9/2 4.79%)"),
       ("聯邦基金利率上限", "2026-09-16", .04, .0375, "加息 25bp（2023 年後首次）", "CNBC / FOMC"),
       ("失業率 (%)", "2026-09", .042, .041, "穩定（峰值 4.6%）", "LongBridge 30771865"),
       ("WTI 原油 ($/bbl)", "2026-09-30", 90.42, 69.50, "Q3 +30%（伊朗衝突）", "FactSet EI"),
       ("金價 Q2 均價 ($/oz)", "2026-Q2", 4520, 3291, "高位", "Webull / FCX 8-K"),
       ("S&P 500 Q3 EPS 增速", "2026-Q3E", .296, .267, "強勁上修", "FactSet EI"),
       ("S&P 500 CY27 EPS 增速", "2027E", .158, None, "放緩", "FactSet EI"),
       ("S&P 500 Fwd P/E", "2026-10-09", 19.1, 20.4, "回落至 10Y 均值", "FactSet EI")]
for i, m in enumerate(MAC):
    r = 6 + i
    put(Mc, r, 1, m[0]); put(Mc, r, 2, m[1], BLUE, al=CENTER)
    fmt = PCT if isinstance(m[2], float) and m[2] < 1 else ('#,##0.00' if m[2] < 200 else '#,##0')
    put(Mc, r, 3, m[2], BLUE, fmt, fl=(GOLD if i == 2 else None)); put(Mc, r, 4, m[3], BLUE, fmt)
    put(Mc, r, 5, m[4], al=WRAP); put(Mc, r, 6, m[5], NOTE, al=WRAP)
    Mc.merge_cells(start_row=r, start_column=6, end_row=r, end_column=10)
Mc["C8"].comment = Comment("關鍵假設：10Y ≈ 4.6%（7/24 為 4.66%、9/2 為 4.79%；10 月實際值未取得）。用於 ERP 與 DCF 無風險利率。", "Analyst")
r = 6 + len(MAC) + 1
section(Mc, r, "B. 當前宏觀體制（Regime）— 方向權重（可調）", 10)
FACT = ["利率上行", "通脹/油價上行", "經濟衰退", "AI 資本開支/用電", "監管政策風險", "美元走強"]
REG = [1.0, 0.8, 0.3, 1.0, 0.5, 0.2]
REGNOTE = ["Fed 9 月加息、點陣圖指向再加一次", "CPI 3.4%、油價 $90", "失業率 4.2% 穩定，概率低", "Q3 IT 盈利 +65%、公用事業大負荷合約", "CCCA、併購審批、關稅", "方向不明"]
header(Mc, r + 1, ["因子"] + FACT, height=30)
put(Mc, r + 2, 1, "體制強度 (0–1)", font(bold=True))
for j, v in enumerate(REG):
    put(Mc, r + 2, 2 + j, v, BLUE, '0.0', al=CENTER, fl=GOLD)
put(Mc, r + 3, 1, "依據", NOTE)
for j, v in enumerate(REGNOTE):
    put(Mc, r + 3, 2 + j, v, NOTE, al=WRAP)
Mc.row_dimensions[r + 3].height = 40
REGROW = r + 2
r = r + 5
section(Mc, r, "C. 個股宏觀敏感度（−2 極負面 … +2 極正面；分析師判斷）→ 宏觀得分 = Σ 敏感度 × 體制強度", 10)
header(Mc, r + 1, ["公司"] + FACT + ["宏觀得分", "韌性分(0–5)", "主要宏觀邏輯"], height=30)
SENS = {"BRK.B": ([1, 0, -1, 0, 0, 0], "浮存金與 $3,600 億現金受惠高利率；分散化抗衰退"),
        "JPM": ([1, 0, -2, 0, 1, 0], "加息擴大 NII（2026 指引 $1,055 億）；衰退時信貸成本上升"),
        "V": ([0, 1, -1, 0, -1, -1], "名義消費額隨通脹上升；CCCA/DOJ 訴訟；強美元壓跨境"),
        "NEE": ([-2, -1, 0, 2, -1, 0], "重資本、高負債最怕加息；Dominion 併購審批；AI 用電長期利好"),
        "CEG": ([-1, 1, -1, 2, 0, 0], "電價上行與數據中心 PPA 利好；併購 Calpine 增加商品價格暴露"),
        "SO": ([-1, 0, 0, 2, 0, 0], "17GW 大型負荷合約鎖定增長；受監管回報抵禦衰退"),
        "WELL": ([-1, 0, 0, 0, 0, 0], "人口老化需求剛性；高估值對利率敏感"),
        "PLD": ([-1, 1, -1, 1, -1, 0], "租金可隨通脹調升；關稅/貿易量風險；數據中心轉型"),
        "EQIX": ([-1, 0, 0, 2, 0, -1], "AI 推理/互聯需求；重資本開支 $50–60 億"),
        "LIN": ([0, 1, -1, 1, 0, -1], "合約通脹轉嫁；電子特氣受惠半導體擴產；工業產量週期"),
        "NEM": ([-1, 2, 1, 0, 0, -2], "實際利率上行不利金價；避險屬性；強美元利空"),
        "FCX": ([0, 1, -2, 2, -1, -1], "電網/數據中心拉動銅需求；衰退敏感；印尼政策風險")}
MACROW = {}
for i, t in enumerate(TICK):
    rr = r + 2 + i; MACROW[t] = rr
    put(Mc, rr, 1, t, font(bold=True))
    for j, v in enumerate(SENS[t][0]):
        put(Mc, rr, 2 + j, v, BLUE, '+0;-0;0', al=CENTER)
    put(Mc, rr, 8, f"=SUMPRODUCT(B{rr}:G{rr},$B${REGROW}:$G${REGROW})", fmt='+0.0;-0.0;0.0', al=CENTER)
    put(Mc, rr, 9, f"=MAX(0,MIN(5,ROUND(2.5+H{rr},0)))", fmt=INT, al=CENTER)
    put(Mc, rr, 10, SENS[t][1], al=WRAP)
    Mc.row_dimensions[rr].height = 30
band(Mc, r + 2, r + 13, 1, 9)
Mc.conditional_formatting.add(f"B{r+2}:G{r+13}", ColorScaleRule(start_type="num", start_value=-2, start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="num", end_value=2, end_color="63BE7B"))
Mc.freeze_panes = "B6"

# ======================= FINANCIALS =======================
Fi = sheet("Financials", "財務數據 FY2025–FY2020（年報）", "來源：LongBridge financial_report（annual）。V 財年截至 9 月。BRK 淨利含未實現投資損益。單位：$B（EPS、ROE 除外）。FCF = OCF − CapEx（有 CapEx 數據者）。",
           [9, 22, 9, 11, 11, 11, 11, 11, 11, 12])
header(Fi, 4, ["Ticker", "指標", "單位", "FY2025", "FY2024", "FY2023", "FY2022", "FY2021", "FY2020", "5Y CAGR/均值"], height=22)
FROW = {}
r = 5
for t in TICK:
    rev, ni, eps, roe, ocf, capex = FIN[t]
    rows = [("營業收入 Revenue", "$B", rev, NUM1), ("淨利潤 Net income", "$B", ni, NUM2), ("每股收益 EPS", "$", eps, NUM2),
            ("ROE", "%", [x / 100 for x in roe], PCT)]
    if ocf: rows.append(("經營現金流 OCF", "$B", ocf, NUM2))
    if capex: rows.append(("資本開支 CapEx", "$B", capex, NUM2))
    FROW[t] = {}
    start = r
    for lab, unit, vals, fmt in rows:
        put(Fi, r, 1, t, font(bold=True)); put(Fi, r, 2, lab); put(Fi, r, 3, unit, al=CENTER)
        for j, v in enumerate(vals):
            put(Fi, r, 4 + j, v, BLUE, fmt)
        key = lab.split()[0]
        FROW[t][key] = r
        if key in ("營業收入", "每股收益"):
            put(Fi, r, 10, f'=IFERROR((D{r}/I{r})^(1/5)-1,"n/m")', fmt=PCT)
        elif key == "ROE":
            put(Fi, r, 10, f"=AVERAGE(D{r}:H{r})", fmt=PCT)
        r += 1
    if ocf and capex:
        put(Fi, r, 1, t, font(bold=True)); put(Fi, r, 2, "自由現金流 FCF"); put(Fi, r, 3, "$B", al=CENTER)
        o, cx = FROW[t]["經營現金流"], FROW[t]["資本開支"]
        for j in range(6):
            col = L(4 + j)
            put(Fi, r, 4 + j, f"={col}{o}-{col}{cx}", fmt=NUM2)
        FROW[t]["FCF"] = r; r += 1
    put(Fi, r, 1, t, font(bold=True)); put(Fi, r, 2, "總資產 / 總負債 FY25"); put(Fi, r, 3, "$B", al=CENTER)
    put(Fi, r, 4, BS[t][0], BLUE, NUM1); put(Fi, r, 5, BS[t][1], BLUE, NUM1)
    put(Fi, r, 6, "負債/資產", al=CENTER); put(Fi, r, 7, f"=E{r}/D{r}", fmt=PCT)
    FROW[t]["BS"] = r; r += 1
    for rr in range(start, r):
        Fi.cell(row=rr, column=1).fill = fill(LIGHT)
    for cc in range(1, 11):
        Fi.cell(row=r - 1, column=cc).border = Border(left=thin, right=thin, top=thin, bottom=Side(style="medium", color=NAVY))
Fi["J5"].comment = Comment("營收/EPS 欄為 FY2020→FY2025 5 年 CAGR；ROE 欄為 FY2021–FY2025 平均。", "Analyst")
Fi.freeze_panes = "D5"

# ======================= COMPS =======================
Co = sheet("Comps", "可比公司估值表 — 4 板塊 × 市值前 3 龍頭", "市場數據：LongBridge calc_indexes / valuation_history / institution_rating（2026-10-09）。Fwd EPS = LongBridge 一致預期中位數（下一財年，≈FY2027E；BRK 為營運 EPS、REIT 為 GAAP EPS）。",
           [8, 20, 9, 16, 10, 11, 9, 9, 10, 9, 8, 8, 9, 9, 9, 9, 9, 9, 8, 9, 9, 9, 8, 10, 9, 9, 8, 7, 7, 7, 7, 7])
cols = ["Ticker", "公司", "板塊", "子行業", "股價 ($)", "市值 ($B)", "股數 (B)", "P/E TTM", "Fwd EPS ($)", "Fwd P/E", "P/B", "股息率",
        "ROE FY25", "ROE 5Y均", "營收 5Y CAGR", "EPS 5Y CAGR", "淨利率 FY25", "負債/資產", "估值指標", "當前", "歷史低", "歷史中位", "歷史高",
        "比歷史 X% 時間便宜", "區間", "行業中位", "YTD", "Buy", "Over", "Hold", "Under+Sell", "總數"]
header(Co, 4, cols, height=40)
CROW = {}
for i, t in enumerate(TICK):
    r = 5 + i; CROW[t] = r
    p, mc, pe, pb, dy, ytd = MKT[t]
    put(Co, r, 1, t, font(bold=True)); put(Co, r, 2, NAME[t]); put(Co, r, 3, SECT[t], al=CENTER); put(Co, r, 4, SUBI[t], al=WRAP)
    put(Co, r, 5, p, BLUE, USD2); put(Co, r, 6, mc, BLUE, NUM1); put(Co, r, 7, f"=F{r}/E{r}", fmt=NUM2)
    put(Co, r, 8, pe, BLUE, MULT); put(Co, r, 9, FWD[t], BLUE, USD2); put(Co, r, 10, f"=E{r}/I{r}", fmt=MULT)
    put(Co, r, 11, pb, BLUE, MULT); put(Co, r, 12, dy / 100, BLUE, PCT)
    fr = FROW[t]
    put(Co, r, 13, f"=Financials!D{fr['ROE']}", GREEN, PCT); put(Co, r, 14, f"=Financials!J{fr['ROE']}", GREEN, PCT)
    put(Co, r, 15, f"=Financials!J{fr['營業收入']}", GREEN, PCT); put(Co, r, 16, f"=Financials!J{fr['每股收益']}", GREEN, PCT)
    put(Co, r, 17, f"=Financials!D{fr['淨利潤']}/Financials!D{fr['營業收入']}", GREEN, PCT)
    put(Co, r, 18, f"=Financials!G{fr['BS']}", GREEN, PCT)
    b = BANDV[t]
    put(Co, r, 19, b[0], al=CENTER)
    for j, v in enumerate(b[1:5]):
        put(Co, r, 20 + j, v, BLUE, MULT)
    put(Co, r, 24, b[5], BLUE, PCT); put(Co, r, 25, b[6], BLUE, al=CENTER); put(Co, r, 26, b[7], BLUE, MULT)
    put(Co, r, 27, ytd / 100, BLUE, PCT)
    ra = RAT[t]
    put(Co, r, 28, ra[3], BLUE, INT, CENTER); put(Co, r, 29, ra[4], BLUE, INT, CENTER); put(Co, r, 30, ra[5], BLUE, INT, CENTER)
    put(Co, r, 31, ra[6] + ra[7], BLUE, INT, CENTER); put(Co, r, 32, ra[8], BLUE, INT, CENTER)
band(Co, 5, 16, 5, 32)
for k, lab in enumerate(["中位數", "平均數"]):
    r = 18 + k
    put(Co, r, 2, lab, font(bold=True))
    fn = "MEDIAN" if k == 0 else "AVERAGE"
    for col in [8, 10, 11, 12, 13, 14, 15, 17, 24]:
        cl = L(col)
        put(Co, r, col, f"={fn}({cl}5:{cl}16)", font(bold=True), PCT if col in (12, 13, 14, 15, 17, 24) else MULT)
put(Co, 21, 1, "註：", font(bold=True), bd=False)
notes = ["BRK、WELL、PLD、EQIX 的估值區間以 P/B 計（LongBridge 提供）；REIT 的 P/E 因折舊失真，見 Valuation 頁 P/FFO。",
         "「比歷史 X% 時間便宜」= 當前倍數低於該區間內多少比例的歷史觀測值（越高越便宜）。NEM 僅 1 年區間、JPM/V/CEG/FCX 為 3 年。",
         "CEG FY2021/2020 EPS 缺失，EPS CAGR 顯示 n/m。WELL/PLD/EQIX 的 GAAP ROE 被折舊壓低，不代表經濟回報。"]
for i, n in enumerate(notes):
    c = Co.cell(row=22 + i, column=1, value=n); c.font = NOTE
Co.conditional_formatting.add("X5:X16", ColorScaleRule(start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=0.5, mid_color="FFFFFF", end_type="num", end_value=1, end_color="63BE7B"))
Co.freeze_panes = "C5"

# ======================= VALUATION =======================
V = sheet("Valuation", "多方法估值 & 反向 DCF", "A1=歷史中位倍數；A2=同業/目標倍數；B=兩階段 DCF/DDM（每股現金流 CF1 第 1–5 年按 g1 增長，之後按 gT 永續）；C=分析師共識目標價。黃底=全局假設。",
           [8, 12, 9, 10, 9, 9, 11, 12, 9, 9, 11, 13, 9, 10, 9, 8, 9, 8, 8, 10, 11, 11, 11, 9, 10, 10, 11, 30])
section(V, 4, "全局假設（可調，全表重算）", 28)
GL = [("無風險利率 rf（=Macro 10Y）", "=Macro!C8", GREEN, PCT), ("股權風險溢價 ERP", 0.045, BLUE, PCT),
      ("權重：A1 歷史倍數", 0.25, BLUE, PCT), ("權重：A2 同業倍數", 0.15, BLUE, PCT), ("權重：B 內在價值", 0.40, BLUE, PCT), ("權重：C 共識目標", 0.20, BLUE, PCT),
      ("低估門檻（安全邊際 ≥）", 0.15, BLUE, PCT), ("高估門檻（安全邊際 ≤）", -0.10, BLUE, PCT)]
for i, (lab, v, f_, fmt) in enumerate(GL):
    r = 5 + i
    put(V, r, 1, lab); V.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    put(V, r, 4, v, f_, fmt, fl=GOLD)
put(V, 13, 1, "權重合計", font(bold=True)); put(V, 13, 4, "=SUM(D7:D10)", font(bold=True), PCT)
RF, ERP, W1, W2, W3, W4, UND, OVR = "$D$5", "$D$6", "$D$7", "$D$8", "$D$9", "$D$10", "$D$11", "$D$12"
hdr = ["Ticker", "A1 指標", "A1 數值", "當前倍數", "歷史中位", "A1 價值", "A2 指標", "A2 數值", "目標倍數", "A2 價值",
       "B 方法", "B 基數 (Fwd)", "轉換/派息率", "CF1 ($)", "Beta", "r", "g1 (1–5年)", "gT", "每股調整", "B 價值",
       "C 共識目標", "加權內在價值", "股價", "安全邊際", "估值判定", "反向DCF 隱含永續g", "隱含g − gT", "方法備註"]
header(V, 15, hdr, height=40)
# A1: (metric label, metric value or formula, hist-median source in Comps col V)
VAL = {
 "BRK.B": ("BVPS", "=Comps!E{r}/Comps!K{r}", "P/B 行業", "=C{v}", 1.48, "營運EPS DCF", 1.00, 0.85, 0.06, 0.03, 77.8,
           "A1/A2 用 P/B 區間；B=2027E 營運 EPS 兩階段 DCF + 超額現金每股 $77.8（($3,655億−$300億回購底線)×50%折讓÷21.57億股）"),
 "JPM": ("EPS TTM", "=Comps!E{r}/Comps!H{r}", "EPS TTM", "=C{v}", 11.99, "可分配盈利 DCF", 0.714, 1.10, 0.05, 0.035, 0,
         "轉換率=1−g/ROE=1−5%/17.5%（留存支持增長）；等同合理 P/B 框架。A2 用行業中位 P/E 11.99x"),
 "V": ("EPS TTM", "=Comps!E{r}/Comps!H{r}", "EPS TTM", "=C{v}", 31.75, "FCF DCF", 1.00, 0.95, 0.11, 0.04, 0,
       "FCF≈淨利（轉換率 >100%）；A2 目標倍數取 MA 當前 P/E 31.75x（行業中位 14x 無可比性）"),
 "NEE": ("EPS TTM", "=Comps!E{r}/Comps!H{r}", "EPS TTM", "=C{v}", 18.73, "DDM", 0.62, 0.80, 0.065, 0.04, 0,
         "DDM：CF1=Fwd EPS×派息率 62%（≈當前 DPS $2.44×1.06）；5Y 中位 25x 含零利率時期，偏高"),
 "CEG": ("EPS TTM", "=Comps!E{r}/Comps!H{r}", "EPS TTM", "=C{v}", 26.74, "FCF DCF", 0.75, 1.20, 0.10, 0.035, 0,
         "A2 目標倍數取 VST 26.7x（同為 IPP）；2026 調整後 EPS 指引 $11–12"),
 "SO": ("EPS TTM", "=Comps!E{r}/Comps!H{r}", "EPS TTM", "=C{v}", 18.73, "DDM", 0.66, 0.75, 0.07, 0.04, 0,
        "DDM：派息率 66%（DPS≈$3.0）；管理層指引 2028 前 EPS 增速 8–9%（取 7% 保守）"),
 "WELL": ("BVPS", "=Comps!E{r}/Comps!K{r}", "FFO 2026E", 6.40, 25.0, "AFFO DCF", 0.90, 0.85, 0.12, 0.04, 0,
          "A2/B 基數用 2026 正常化 FFO 指引中值 $6.40（8-K 7/27）；AFFO≈90% FFO；目標 P/FFO 25x 為分析師估計"),
 "PLD": ("BVPS", "=Comps!E{r}/Comps!K{r}", "Core FFO 2026E", 6.26, 22.0, "AFFO DCF", 0.85, 1.05, 0.07, 0.035, 0,
         "Core FFO 指引中值 $6.26（$6.22–6.30）；AFFO≈85%；目標 P/FFO 22x（物流 REIT 歷史中樞，估計）"),
 "EQIX": ("BVPS", "=Comps!E{r}/Comps!K{r}", "AFFO 2026E", 42.99, 24.0, "AFFO DCF", 1.00, 0.90, 0.09, 0.04, 0,
          "AFFO 指引中值 $42.99（$42.69–43.29）；目標 P/AFFO 24x（估計）"),
 "LIN": ("EPS TTM", "=Comps!E{r}/Comps!H{r}", "EPS TTM", "=C{v}", 29.45, "FCF DCF", 0.85, 0.90, 0.08, 0.035, 0,
         "2026 調整後 EPS 指引 $17.70–17.90（+8–9%）；FCF 轉換率 85%（增長型資本開支）"),
 "NEM": ("EPS TTM", "=Comps!E{r}/Comps!H{r}", "EPS TTM", "=C{v}", 8.53, "正常化 DCF", 0.90, 0.90, -0.08, 0.02, 0,
         "g1=−8% 反映金價由 ~$4,500 向中週期回歸假設；歷史區間僅 1 年（高金價期）"),
 "FCX": ("EPS TTM", "=Comps!E{r}/Comps!H{r}", "EPS TTM", "=C{v}", 31.20, "正常化 DCF", 0.70, 1.50, 0.04, 0.025, 0,
         "Grasberg 2027 年底恢復滿產支撐量增，但銅價 $6.5/lb 處高位；Beta 1.5"),
}
VROW = {}
for i, t in enumerate(TICK):
    r = 16 + i; VROW[t] = r; cr = CROW[t]
    a1lab, a1f, a2lab, a2v, tgtm, bm, conv, beta, g1, gT, adj, note = VAL[t]
    put(V, r, 1, t, font(bold=True))
    put(V, r, 2, a1lab, al=CENTER)
    put(V, r, 3, a1f.format(r=cr), GREEN, USD2)
    put(V, r, 4, f"=W{r}/C{r}", fmt=MULT)
    put(V, r, 5, f"=Comps!V{cr}", GREEN, MULT)
    put(V, r, 6, f"=C{r}*E{r}", fmt=USD2)
    put(V, r, 7, a2lab, al=CENTER)
    if isinstance(a2v, str):
        put(V, r, 8, a2v.format(v=r), fmt=USD2)
    else:
        put(V, r, 8, a2v, BLUE, USD2)
    put(V, r, 9, tgtm, BLUE, MULT, fl=GOLD)
    put(V, r, 10, f"=H{r}*I{r}", fmt=USD2)
    put(V, r, 11, bm, al=CENTER)
    if t in ("WELL", "PLD", "EQIX"):
        put(V, r, 12, f"=H{r}", fmt=USD2)
    else:
        put(V, r, 12, f"=Comps!I{cr}", GREEN, USD2)
    put(V, r, 13, conv, BLUE, PCT, fl=GOLD)
    put(V, r, 14, f"=L{r}*M{r}", fmt=USD2)
    put(V, r, 15, beta, BLUE, '0.00')
    put(V, r, 16, f"={RF}+O{r}*{ERP}", fmt=PCT)
    put(V, r, 17, g1, BLUE, PCT, fl=GOLD); put(V, r, 18, gT, BLUE, PCT, fl=GOLD)
    put(V, r, 19, adj, BLUE, USD2)
    # two-stage: CF1/(r-g1)*(1-((1+g1)/(1+r))^5) + CF1*(1+g1)^4*(1+gT)/(r-gT)/(1+r)^5
    put(V, r, 20, f"=IF(ABS(P{r}-Q{r})<0.0001,N{r}*5/(1+P{r}),N{r}/(P{r}-Q{r})*(1-((1+Q{r})/(1+P{r}))^5))"
                  f"+N{r}*(1+Q{r})^4*(1+R{r})/(P{r}-R{r})/(1+P{r})^5+S{r}", fmt=USD2)
    put(V, r, 21, f"=Comps!X{cr}" if False else RAT[t][0], BLUE, USD2)
    put(V, r, 22, f"=F{r}*{W1}+J{r}*{W2}+T{r}*{W3}+U{r}*{W4}", font(bold=True), USD2, fl=LIGHT)
    put(V, r, 23, f"=Comps!E{cr}", GREEN, USD2)
    put(V, r, 24, f"=V{r}/W{r}-1", font(bold=True), PCT)
    put(V, r, 25, f'=IF(X{r}>={UND},"低估",IF(X{r}<={OVR},"高估","合理"))', font(bold=True), al=CENTER)
    put(V, r, 26, f"=P{r}-N{r}*(1+Q{r})/W{r}", fmt=PCT)
    put(V, r, 27, f"=Z{r}-R{r}", fmt=PCT)
    put(V, r, 28, note, NOTE, al=WRAP)
    V.row_dimensions[r].height = 40
band(V, 16, 27, 1, 27)
V.conditional_formatting.add("X16:X27", ColorScaleRule(start_type="num", start_value=-0.3, start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="num", end_value=0.3, end_color="63BE7B"))
V.conditional_formatting.add("Y16:Y27", CellIsRule(operator="equal", formula=['"低估"'], fill=fill("C6EFCE"), font=font("006100", True)))
V.conditional_formatting.add("Y16:Y27", CellIsRule(operator="equal", formula=['"高估"'], fill=fill("FFC7CE"), font=font("9C0006", True)))
r = 29
section(V, r, "讀表說明", 28)
expl = ["• 反向 DCF：在當前股價下，市場隱含的永續增長率 g = r − CF2/P（CF2 = CF1×(1+g1)，單階段 Gordon 反解）。隱含 g 遠高於 gT（第 AA 欄 > +1.5%）代表市場已定價高增長，容錯空間小。",
        "• 共識目標價（C）只佔 20%：分析師目標價有樂觀與追漲偏差，僅作情緒參考。",
        "• REIT 的 B 值以 AFFO 為現金流基數；公用事業用 DDM（股息為股東實得現金，重資本開支令 FCF 為負）；銀行用可分配盈利（=EPS×(1−g/ROE)）。",
        "• 所有 g1/gT/轉換率/Beta/目標倍數為分析師假設（藍字黃底），可直接修改測試敏感度。Beta 為 Blume 調整後估計值（0.67×原始+0.33）；公用事業原始 Beta 偏低會低估其利率風險。"]
for i, e in enumerate(expl):
    c = V.cell(row=r + 1 + i, column=1, value=e); c.font = BLACK
    V.merge_cells(start_row=r + 1 + i, start_column=1, end_row=r + 1 + i, end_column=27)
# sensitivity table for picks: B value vs r and g1 (LIN)
r = 35
section(V, r, "敏感度：LIN 內在價值（B）對 折現率 r × 第一階段增長 g1（gT 取表內值）", 28)
put(V, r + 1, 1, "r ↓ / g1 →", font(bold=True))
gs = [0.06, 0.07, 0.08, 0.09, 0.10]; rs = [0.075, 0.080, 0.085, 0.090, 0.095]
lr = VROW["LIN"]
for j, g in enumerate(gs):
    put(V, r + 1, 2 + j, g, BLUE, PCT, al=CENTER, fl=LIGHT)
for i, rv in enumerate(rs):
    rr = r + 2 + i
    put(V, rr, 1, rv, BLUE, PCT, fl=LIGHT)
    for j in range(len(gs)):
        gc = f"{L(2+j)}${r+1}"; rc = f"$A{rr}"
        put(V, rr, 2 + j, f"=IF(ABS({rc}-{gc})<0.0001,$N${lr}*5/(1+{rc}),$N${lr}/({rc}-{gc})*(1-((1+{gc})/(1+{rc}))^5))+$N${lr}*(1+{gc})^4*(1+$R${lr})/({rc}-$R${lr})/(1+{rc})^5", fmt=USD0, al=CENTER)
put(V, r + 1, 8, "現價", font(bold=True)); put(V, r + 2, 8, f"=W{lr}", GREEN, USD2)
V.freeze_panes = "B16"

# ======================= SCORECARD =======================
Sc = sheet("Scorecard", "巴菲特準則評分卡（質量）+ 估值分", "0–5 分；藍字=分析師判斷（理由見右欄），黑字=公式（宏觀韌性連結 Macro，估值分由安全邊際與 5 年預期回報計算）。",
           [8, 9, 9, 9, 9, 9, 9, 10, 10, 10, 10, 10, 12, 30, 44, 44, 40])
section(Sc, 4, "權重（質量分內部，可調）", 17)
CR = ["業務易懂(能力圈)", "護城河寬度", "長期 ROE/ROIC 質量", "資產負債表穩健", "管理層（誠信+資本配置）", "宏觀韌性"]
WT = [0.10, 0.30, 0.15, 0.10, 0.25, 0.10]
for j, (c_, w) in enumerate(zip(CR, WT)):
    put(Sc, 5, 2 + j, c_, font(bold=True), al=CENTER); put(Sc, 6, 2 + j, w, BLUE, PCT, al=CENTER, fl=GOLD)
Sc.row_dimensions[5].height = 30
put(Sc, 6, 1, "權重", font(bold=True)); put(Sc, 6, 8, "=SUM(B6:G6)", font(bold=True), PCT)
put(Sc, 7, 1, "綜合分 = 質量分 × 質量權重 + 估值分 × (1−質量權重)；質量權重：", NOTE, bd=False)
put(Sc, 7, 9, 0.6, BLUE, PCT, fl=GOLD)
QW = "$I$7"
header(Sc, 9, ["Ticker"] + CR + ["質量分 (0–100)", "安全邊際", "5Y 預期年化", "估值分 (0–100)", "綜合分", "排名", "護城河類型 / 寬度", "管理層評估", "巴菲特視角"], height=40)
SCORE = {
 "BRK.B": ([4, 5, 3, 5, 4], "寬：低成本浮存金、永久資本、分散化現金流、「首選買家」聲譽",
           "Greg Abel 2026/1 接任 CEO；Q2 營運盈利 +16% 至 $129.8 億；回購重啟（Q2 $45 億，判斷內在價值高於市價）；$100 億投資 Alphabet、$68 億收購 Taylor Morrison。後巴菲特時代資本配置能力待驗證",
           "巴菲特親手打造；P/B 1.49 接近其歷史回購區間；$3,655 億現金在加息期是期權，不是負擔"),
 "JPM": ([4, 4, 4, 3, 5], "寬：規模、低成本存款、全業務線、科技投入",
         "Jamie Dimon 2005 年起任 CEO，ROE 16–18%、堡壘式資產負債表；2026 NII 指引上調至 ~$1,055 億；接班「數年內」、兩位聯席總裁",
         "巴菲特曾稱 Dimon 為最佳銀行家；但 P/B 2.5x 高於行業 1.57x，價格已反映質量"),
 "V": ([5, 5, 5, 4, 4], "極寬：雙邊網絡效應 + 雙寡頭；毛利 97.8%、ROE 53%",
       "Ryan McInerney 2023 年起；每季回饋股東 ~$61 億（回購+股息）；需處理 CCCA、DOJ 反壟斷、交換費訴訟（每季撥備 $6–10 億）",
         "巴菲特持有 V/MA；典型「以合理價格買入卓越公司」，但 33x TTM 不便宜"),
 "NEE": ([3, 3, 3, 2, 3], "窄：FPL 受監管壟斷 + 再生能源開發規模",
         "John Ketchum 2022 年起；2026/5 宣布 $668 億收購 Dominion，溢價受質疑、股價半年 −17%；FCF 長期為負依賴融資",
         "巴菲特經 BHE 持有公用事業，但偏好低槓桿；NEE 估值處 5 年低位（比 87% 時間便宜）是逆向機會而非質量標的"),
 "CEG": ([3, 3, 3, 3, 3], "窄–寬：核電機組不可複製",
         "Joe Dominguez；Calpine 併購整合中；2026 指引 $11–12 低於預期、缺新超大規模 PPA",
         "盈利受電價與合約時點驅動，可預測性不足"),
 "SO": ([4, 3, 3, 2, 3], "窄：東南部受監管壟斷",
        "Chris Womack 2023 年起；Vogtle 3/4 超支延誤為歷史污點；$810 億資本計劃、17GW 大負荷合約",
        "業務易懂、股息穩，但高負債（負債/資產 75%）與監管風險"),
 "WELL": ([3, 3, 2, 4, 4], "窄：長者住宅營運平台 + 數據化管理",
          "Shankh Mitra 2020 年起；FFO 指引上調至 $6.36–6.44、股息 +15%；資本配置出色",
          "好公司但 P/FFO 35x、P/B 處 5 年高位（僅比 8% 時間便宜），無安全邊際"),
 "PLD": ([4, 4, 3, 4, 5], "寬：全球物流網絡規模 + 客戶關係 + 低資本成本",
         "Hamid Moghadam（創辦人，執行主席）/ Dan Letter（CEO）；Core FFO 指引上調至 $6.22–6.30；數據中心電力管線 5.8GW；擬收購 SEGRO 待觀察",
         "類似「收費公路」的物流地產；P/FFO 20.7x 處合理區間，股息 3.3%"),
 "EQIX": ([3, 4, 3, 3, 3], "寬：互聯生態網絡效應",
          "Adaire Fox-Martin 2024 年起；AFFO 指引上調至 $42.69–43.29；2024 年遭沽空報告質疑 AFFO 口徑；資本開支 $50–60 億",
          "護城河強但資本密集、P/B 處 5 年高位"),
 "LIN": ([5, 5, 4, 4, 5], "寬：三寡頭 + 長約 + 區域密度 + 轉換成本",
         "Sanjiv Lamba 2022 年起；EPS 指引 +8–9%、backlog $110 億；ROE 由 5% 升至 18%；紀律性回購",
         "最接近「工業版收費站」；P/E 處 5 年最低 9%（比 91% 時間便宜），質量+估值兼得"),
 "NEM": ([3, 1, 2, 4, 3], "無：價格接受者",
         "Natascha Viljoen（2026 起任 CEO）；Q2 FCF 創紀錄 $22 億、淨現金 $34 億、$60 億回購；Newcrest 併購攤薄歷史",
         "巴菲特曾形容黃金「不產生任何東西」；盈利完全取決於金價"),
 "FCX": ([3, 2, 3, 3, 3], "窄：一級資產但無定價權",
         "Kathleen Quirk 2024 年起；Grasberg 2025 年泥石流事故後產能 65%，2027 年底恢復；Q3 實現銅價 >$6.5/lb",
         "週期性、資本密集；不符巴菲特偏好"),
}
SROW = {}
for i, t in enumerate(TICK):
    r = 10 + i; SROW[t] = r
    s5, moat, mgmt, buff = SCORE[t]
    put(Sc, r, 1, t, font(bold=True))
    for j, v in enumerate(s5):
        put(Sc, r, 2 + j, v, BLUE, INT, al=CENTER)
    put(Sc, r, 7, f"=Macro!I{MACROW[t]}", GREEN, INT, al=CENTER)
    put(Sc, r, 8, f"=SUMPRODUCT(B{r}:G{r},$B$6:$G$6)/5*100", font(bold=True), '0', al=CENTER)
    put(Sc, r, 9, f"=Valuation!X{VROW[t]}", GREEN, PCT)
    put(Sc, r, 10, f"=Picks!N{6+i}", GREEN, PCT)
    put(Sc, r, 11, f"=ROUND(0.5*MAX(0,MIN(100,(I{r}+0.3)/0.6*100))+0.5*MAX(0,MIN(100,(J{r}-0.02)/0.10*100)),0)", fmt='0', al=CENTER)
    put(Sc, r, 12, f"=ROUND(H{r}*{QW}+K{r}*(1-{QW}),1)", font(bold=True), '0.0', al=CENTER, fl=LIGHT)
    put(Sc, r, 13, f"=RANK(L{r},$L$10:$L$21)", font(bold=True), INT, al=CENTER)
    put(Sc, r, 14, moat, al=WRAP); put(Sc, r, 15, mgmt, al=WRAP); put(Sc, r, 16, buff, al=WRAP)
    Sc.row_dimensions[r].height = 64
band(Sc, 10, 21, 1, 13)
Sc.conditional_formatting.add("H10:H21", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
Sc.conditional_formatting.add("L10:L21", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
put(Sc, 23, 1, "估值分 = 50% × 安全邊際分（−30%→0 分，+30%→100 分）+ 50% × 5 年預期年化回報分（2%→0 分，12%→100 分）。", NOTE, bd=False)
r = 25
section(Sc, r, "巴菲特選股邏輯（本評分卡的框架來源）", 17)
BUF = ["1. 能力圈：只買自己看得懂、10 年後大概率仍存在的業務（→ 業務易懂分）。",
       "2. 持久競爭優勢（Moat）：「我們尋找被寬闊護城河保護的經濟城堡」（→ 權重最高 30%）。",
       "3. 高且穩定的 ROE、低負債：不靠槓桿也能賺高回報（→ ROE 質量、資產負債表）。",
       "4. 誠實能幹、以股東利益為先的管理層，資本配置能力決定長期複利（→ 25%）。",
       "5. 安全邊際：「以合理價格買入卓越公司，勝過以低價買入平庸公司」（→ 估值分，佔綜合分 40%）。",
       "6. 避開：需持續大量資本開支且無定價權的週期業務（礦業）；價格完全取決於商品的資產。"]
for i, b in enumerate(BUF):
    c = Sc.cell(row=r + 1 + i, column=1, value=b); c.font = BLACK
    Sc.merge_cells(start_row=r + 1 + i, start_column=1, end_row=r + 1 + i, end_column=16)
Sc.freeze_panes = "B10"

# ======================= PICKS =======================
P = sheet("Picks", "5 年期選股 — 預期回報、排序與組合", "5Y 預期年化 = [(1+g)^5 × 退出倍數/買入倍數]^(1/5) − 1 + 股息率。g、退出倍數為分析師假設（黃底）；綜合分來自 Scorecard。",
          [8, 20, 9, 11, 12, 10, 11, 10, 10, 9, 9, 9, 9, 11, 10, 8, 10, 12, 50, 50])
header(P, 5, ["Ticker", "公司", "板塊", "股價", "倍數基數", "基數值", "買入倍數", "退出倍數", "g (5年CAGR)", "股息率", "倍數變化/年", "盈利貢獻/年", "—",
              "5Y 預期年化", "綜合分", "排名", "建議", "組合權重", "買入理由（為什麼選它）", "出場/判斷錯誤條件（Kill criteria）"], height=40)
PK = {
 "BRK.B": ("BVPS", 1.46, 0.08, "核心持有", 0.225, "質量最高、下行保護最強：$3,655 億現金在加息期賺取 ~4% 利息並等待危機中的收購機會；P/B 1.49 接近歷史回購區間，Abel 已用真金白銀回購（Q2 $45 億）；營運盈利 +16%。是組合的「壓艙石」。",
           "回購在 P/B<1.4 時停止（反映管理層不再認為便宜）；高價大型併購；保險承保結構性虧損；P/B>1.8 減持"),
 "JPM": ("Fwd EPS", 13.0, 0.07, "核心持有（逢低加碼）", 0.125, "美國最佳銀行：ROE 17–18%、CET1 14%+；加息擴大 NII（2026 指引 ~$1,055 億）；估值處 3 年中位附近，不貴不便宜，建議股價 <$300（≈2.25x P/B）時加碼。",
         "CET1 < 12.5%；淨撇賬率 > 1.5%；接班混亂；P/TBV > 3x 減持"),
 "V": ("Fwd EPS", 25.0, 0.11, "核心持有", 0.20, "全美最寬護城河之一（雙寡頭網絡效應、毛利 98%、ROE 53%）；EPS 5 年 CAGR 16%；Fwd P/E 26.6x 處 3 年區間中段，屬「合理價格買卓越公司」；穩定幣以結算層納入而非被取代（年化 $46 億）。",
       "CCCA 通過或美國交換費上限；支付量增速 < 5% 連續 4 季；穩定幣在美國零售端實質分流；Fwd P/E > 32x 減持"),
 "NEE": ("Fwd EPS", 20.0, 0.07, "逆向觀察（小倉位）", 0.075, "估值處 5 年低位（P/E 17.4x，比 87% 時間便宜）、股息率 3.2%；AI 數據中心用電長期利好；若 10/21 Q3 業績與 Dominion 審批順利，估值修復空間最大（5Y 預期回報最高）。",
         "Dominion 審批附帶苛刻條件或需大量增發股份；10Y > 5.25%；信用評級下調"),
 "CEG": ("Fwd EPS", 22.0, 0.10, "觀察", 0, "核電稀缺資產，但盈利可預測性不足、估值仍在 3 年中位上方。", "—"),
 "SO": ("Fwd EPS", 18.5, 0.07, "備選", 0, "17GW 大負荷合約鎖定增長、股息 3.5%；但高負債、共識多為 Hold；作公用事業備選。", "—"),
 "WELL": ("FFO 2026E", 26.0, 0.12, "不買（估值過高）", 0, "營運極佳但 P/FFO 35x、P/B 處 5 年高位，估值壓縮風險抵消增長。", "—"),
 "PLD": ("FFO 2026E", 21.0, 0.07, "核心持有", 0.15, "全球物流地產龍頭：P/FFO 20.7x 合理、股息 3.3%、Core FFO 指引上調；數據中心電力管線 5.8GW 提供期權；管理層資本配置業內最佳。是低估房地產板塊中「質量+價格」最平衡的一檔。",
         "出租率 < 93%；Core FFO 增速 < 3%；SEGRO 收購溢價過高；P/FFO > 27x 減持"),
 "EQIX": ("FFO 2026E", 23.0, 0.09, "候補（第 6 名）", 0, "互聯護城河強、AI 推理需求受惠，但資本密集、P/B 處 5 年高位；回調 10–15% 可替補 JPM/PLD。", "—"),
 "LIN": ("Fwd EPS", 27.0, 0.085, "核心持有（首選）", 0.225, "質量與估值兼得的首選：工業氣體三寡頭、15–20 年照付不議合約、通脹可轉嫁；ROE 由 5% 升至 18%；P/E 30.8x 處 5 年最低區域（比 91% 時間便宜）；EPS 指引 +8–9%、backlog $110 億（電子特氣受惠半導體擴產）。",
         "Backlog < $80 億；EPS 增速連續 2 年 < 5%；ROE 掉回 15% 以下；P/E > 35x 減持"),
 "NEM": ("Fwd EPS", 12.0, -0.05, "不買（週期頂部）", 0, "盈利取決於金價 ~$4,500；P/E 低是週期頂部特徵，不符合巴菲特準則。", "—"),
 "FCX": ("Fwd EPS", 17.0, 0.04, "不買", 0, "銅長期需求好，但銅價高位 + Grasberg 恢復風險 + 無定價權。", "—"),
}
for i, t in enumerate(TICK):
    r = 6 + i; cr = CROW[t]; vr = VROW[t]
    base, exitm, g, rec, w, why, kill = PK[t]
    put(P, r, 1, t, font(bold=True)); put(P, r, 2, NAME[t]); put(P, r, 3, SECT[t], al=CENTER)
    put(P, r, 4, f"=Comps!E{cr}", GREEN, USD2)
    put(P, r, 5, base, al=CENTER)
    if base == "BVPS":
        put(P, r, 6, f"=Valuation!C{vr}", GREEN, USD2)
    elif base.startswith("FFO"):
        put(P, r, 6, f"=Valuation!H{vr}", GREEN, USD2)
    else:
        put(P, r, 6, f"=Comps!I{cr}", GREEN, USD2)
    put(P, r, 7, f"=D{r}/F{r}", fmt=MULT)
    put(P, r, 8, exitm, BLUE, MULT, fl=GOLD)
    put(P, r, 9, g, BLUE, PCT, fl=GOLD)
    put(P, r, 10, f"=Comps!L{cr}" if t != "WELL" else 0.015, GREEN if t != "WELL" else BLUE, PCT)
    put(P, r, 11, f"=(H{r}/G{r})^(1/5)-1", fmt=PCT)
    put(P, r, 12, f"=I{r}", fmt=PCT)
    put(P, r, 13, "", bd=False)
    put(P, r, 14, f"=((1+I{r})^5*H{r}/G{r})^(1/5)-1+J{r}", font(bold=True), PCT, fl=LIGHT)
    put(P, r, 15, f"=Scorecard!L{SROW[t]}", GREEN, '0.0', al=CENTER)
    put(P, r, 16, f"=Scorecard!M{SROW[t]}", GREEN, INT, al=CENTER)
    put(P, r, 17, rec, font(bold=True), al=CENTER)
    put(P, r, 18, w, BLUE, PCT, al=CENTER, fl=(GOLD if w > 0 else None))
    put(P, r, 19, why, al=WRAP); put(P, r, 20, kill, al=WRAP)
    P.row_dimensions[r].height = 70 if w > 0 else 34
P["J12"].comment = Comment("WELL：按新季度股息 $0.85×4 ÷ 股價 ≈ 1.5%（TTM 1.36% 未反映 15% 加息）。", "Analyst")
band(P, 6, 17, 1, 16)
P.conditional_formatting.add("N6:N17", ColorScaleRule(start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=0.08, mid_color="FFEB84", end_type="num", end_value=0.12, end_color="63BE7B"))
put(P, 19, 2, "組合合計", font(bold=True)); put(P, 19, 18, "=SUM(R6:R17)", font(bold=True), PCT, al=CENTER)
put(P, 20, 2, "組合加權 5Y 預期年化", font(bold=True)); put(P, 20, 14, "=SUMPRODUCT(N6:N17,R6:R17)/R19", font(bold=True), PCT, fl=GOLD)
put(P, 21, 2, "組合加權股息率", font(bold=True)); put(P, 21, 14, "=SUMPRODUCT(J6:J17,R6:R17)/R19", font(bold=True), PCT)
put(P, 22, 2, "組合加權綜合分", font(bold=True)); put(P, 22, 14, "=SUMPRODUCT(O6:O17,R6:R17)/R19", font(bold=True), '0.0')
put(P, 23, 2, "對照：S&P 500 隱含回報（盈利收益率+CY27增速近似）", NOTE, bd=False)
put(P, 23, 14, "=1/Sector_Map!C9", GREEN, PCT)
P["N23"].comment = Comment("S&P 500 Fwd 盈利收益率 = 1/19.1 ≈ 5.2%（實質長期回報近似值，未計增長與估值變化）。", "Analyst")
put(P, 3, 1, "結論：5 年期首選 LIN、BRK.B、V、PLD、JPM 為核心，NEE 作逆向小倉位。核心特徵＝寬護城河 + 優秀管理層 + 估值不高於歷史中位；刻意迴避週期頂部（NEM/FCX/能源）與估值過高的好公司（WELL/EQIX）。",
    font(NAVY, True, 10), bd=False)
P.merge_cells("A3:T3")
P.freeze_panes = "C6"

# ======================= SUMMARY =======================
Su = sheet("Summary", "結論摘要 Executive Summary", "數據截至 2026-10-09；所有數字連結至對應工作表，可隨假設重算。",
           [4, 26, 14, 14, 14, 14, 14, 14, 50])
section(Su, 4, "一、核心結論", 9)
KEY = ["1) 大盤：S&P 500 Fwd P/E 19.1x = 10Y 均值，但 Trailing 26.0x 偏高；盈利動能極強（Q3E +29.6%、CY26E +32.7%），CY27 放緩至 +15.8%。整體「合理偏貴」，選股重於擇時。",
       "2) 板塊：相對 10Y 均值最低估——公用事業(−10.6%)、房地產(−10.2%)、原材料(−6.7%)；最高估——能源(+22.7%，盈利頂部陷阱)、工業(+14%)、醫療(+10%，LLY 拉高)。",
       "3) 宏觀：伊朗衝突推高油價(Q3 均價 $85.7)與 CPI(3.4%)，Fed 9 月加息至 3.75–4.00%。利率敏感板塊（公用、REIT）被錯殺，但「高利率更久」仍是主要風險。",
       "4) 估值方法：板塊 P/E 只能做首輪篩選。REIT 要看 P/FFO（WELL GAAP P/E 105x vs P/FFO 35x）、銀行看 P/B-ROE、礦業看中週期價格，否則會誤判。",
       "5) 個股：12 家龍頭中，LIN 是唯一「質量頂級 + 倍數處 5 年低位」的組合（但 DCF 顯示市場已隱含 ~5% 永續增長，屬合理價而非深度價值）；BRK.B、V、JPM、PLD 屬「合理價格買卓越公司」；NEE 最便宜但風險最高；WELL/EQIX 好公司但太貴；NEM/FCX 處週期頂部。",
       "6) 估值結論：12 家中沒有一家達到「安全邊際 ≥15%」的低估門檻（NEE 最接近，+12%）；優質龍頭（LIN/V/BRK/PLD/JPM）都在「合理」區間——這是一個要為質量付合理價的市場。5 年組合加權預期年化 ≈ 9–10%（下表公式連動），高於 S&P 500 盈利收益率 5.2%。"]
for i, k in enumerate(KEY):
    c = Su.cell(row=5 + i, column=2, value=k); c.font = BLACK; c.alignment = WRAP
    Su.merge_cells(start_row=5 + i, start_column=2, end_row=5 + i, end_column=9); Su.row_dimensions[5 + i].height = 46
r = 12
section(Su, r, "二、最終 5 年期組合", 9)
header(Su, r + 1, ["#", "公司", "Ticker", "建議", "組合權重", "股價", "加權內在價值", "5Y 預期年化", "一句話理由"], height=30)
FINAL = [("LIN", "質量+估值兼得，P/E 處 5 年最低 9% 區域"), ("BRK.B", "壓艙石：$3,655 億現金 + 回購重啟 + 加息受惠"),
         ("V", "最寬護城河之一，價格合理"), ("PLD", "低估房地產中質量/價格最平衡，股息 3.3%"),
         ("JPM", "最佳銀行，加息受惠；<$300 加碼"), ("NEE", "逆向：5 年估值低位，等待 Q3 與 Dominion 審批")]
for i, (t, why) in enumerate(FINAL):
    rr = r + 2 + i; pr = 6 + TICK.index(t); vr = VROW[t]
    put(Su, rr, 1, i + 1, al=CENTER); put(Su, rr, 2, NAME[t], font(bold=True), al=Alignment(vertical='center')); put(Su, rr, 3, t, al=CENTER); Su.row_dimensions[rr].height = 30
    put(Su, rr, 4, f"=Picks!Q{pr}", GREEN, al=CENTER); put(Su, rr, 5, f"=Picks!R{pr}", GREEN, PCT, al=CENTER)
    put(Su, rr, 6, f"=Picks!D{pr}", GREEN, USD2); put(Su, rr, 7, f"=Valuation!V{vr}", GREEN, USD2)
    put(Su, rr, 8, f"=Picks!N{pr}", GREEN, PCT); put(Su, rr, 9, why, al=WRAP)
rr = r + 8
put(Su, rr, 2, "組合合計 / 加權預期年化", font(bold=True)); put(Su, rr, 5, "=SUM(E14:E19)", font(bold=True), PCT, al=CENTER)
put(Su, rr, 8, "=Picks!N20", font(bold=True), PCT, fl=GOLD)
r = rr + 2
section(Su, r, "三、刻意不選的龍頭與原因", 9)
NOT = [("WELL / EQIX", "好公司、壞價格：P/B 均處 5 年高位（僅比 8–9% 時間便宜），估值壓縮會吃掉增長。"),
       ("NEM / FCX", "價格接受者、盈利處金/銅價高位；低 P/E 是週期頂部訊號，不符合巴菲特準則。"),
       ("CEG / SO", "CEG 盈利可預測性不足；SO 高負債、共識多為 Hold，作公用事業備選。"),
       ("能源板塊", "Fwd P/E 13.5x 看似最低，但高於 10Y 均值 23%，CY27 EPS −9.3%，油價含地緣溢價。")]
for i, (a, b) in enumerate(NOT):
    put(Su, r + 1 + i, 2, a, font(bold=True), al=Alignment(vertical='center')); put(Su, rr_ := r + 1 + i, 3, b, al=Alignment(wrap_text=True, vertical='center'))
    Su.merge_cells(start_row=r + 1 + i, start_column=3, end_row=r + 1 + i, end_column=9); Su.row_dimensions[r + 1 + i].height = 28
r = r + 6
section(Su, r, "四、未來 3 個月關鍵催化劑", 9)
CAT = ["10/14 美國 9 月 CPI（預期 3.7%）→ 決定 Fed 是否再加息，直接影響 NEE/SO/PLD/WELL",
       "10 月中 JPM Q3 業績（NII 指引）、10/21 NEE Q3 業績（Dominion 進展）、10 月底 V 財年業績",
       "11–12 月 FOMC：點陣圖指向年底 4.00–4.25%；油價與伊朗局勢"]
for i, c_ in enumerate(CAT):
    cc = Su.cell(row=r + 1 + i, column=2, value="• " + c_); cc.font = BLACK
    Su.merge_cells(start_row=r + 1 + i, start_column=2, end_row=r + 1 + i, end_column=9)
put(Su, r + 5, 2, "免責聲明：僅供研究參考，不構成投資建議。", NOTE, bd=False)

# ======================= SOURCES =======================
So = sheet("Sources", "數據來源、假設與免責聲明", "所有硬編碼數字的出處。", [4, 40, 90])
header(So, 4, ["#", "數據", "來源"], height=22)
SRC = [("板塊 Fwd P/E、5Y/10Y 均值、CY26/27 盈利/營收增速、目標價、評級分佈", "FactSet Earnings Insight, 2026-10-09 (pp.17, 29, 30, 34, 37) https://advantage.factset.com/hubfs/Website/Resources%20Section/Research%20Desk/Earnings%20Insight/EarningsInsight_100926.pdf"),
       ("股價、市值、P/E、P/B、股息率、YTD", "LongBridge OpenAPI calc_indexes（2026-10-09 收盤）"),
       ("估值歷史區間與分位、行業中位數", "LongBridge OpenAPI valuation_history / industry_valuation_dist"),
       ("FY2020–FY2025 財務數據", "LongBridge OpenAPI financial_report（annual）"),
       ("一致預期 EPS、分析師目標價與評級", "LongBridge OpenAPI forecast_eps / institution_rating"),
       ("CPI、失業率", "LongBridge OpenAPI macrodata（30771871, 30771865）"),
       ("Fed 2026-09-16 加息至 3.75–4.00%", "https://www.cnbc.com/2026/09/16/fed-rate-decision-september-2026.html"),
       ("10Y 殖利率 4.66%（7/24）、4.79%（9/2）", "https://ycharts.com/indicators/10_year_treasury_rate_weekly（10 月實際值未取得，假設 4.6%）"),
       ("2026 通脹由能源/伊朗衝突驅動、5 月 CPI 4.2%", "https://cryptobriefing.com/us-inflation-rises-oil-prices-falling/ ；https://www.idnfinancials.com/news/63751/iran-war-pushes-us-inflation-to-3-8-trump-under-pressure"),
       ("BRK Q2 2026 營運盈利、回購、現金、Alphabet、Taylor Morrison", "https://www.cnbc.com/2026/08/08/berkshire-hathaway-earnings-q2-2026.html ；SEC 10-Q"),
       ("JPM 2026 NII 指引、Dimon 接班", "https://www.marketbeat.com/instant-alerts/jpmorgan-chase-co-q2-earnings-call-highlights-2026-07-14/ ；https://www.americanbanker.com/news/dimon-says-jpmorgan-ceo-transition-remains-years-away"),
       ("V 穩定幣、CCCA、交換費訴訟", "https://www.nasdaq.com/articles/time-buy-visa-its-stablecoin-strategy-redefines-payments ；https://247wallst.com/investing/2026/07/13/visas-700-billion-payments-empire-faces-new-disruption-risks/"),
       ("NEE 收購 Dominion、股價下跌原因", "https://www.fool.com/investing/2026/05/18/why-nextera-energy-stock-fell-54-today/ ；https://www.gurufocus.com/news/8910473/"),
       ("CEG 2026 指引 $11–12", "https://www.tikr.com/blog/constellation-energy-stock-is-down-13-in-2026-what-the-q1-2026-earnings-must-show"),
       ("SO 指引、17GW 大負荷、$810 億資本計劃", "https://www.nasdaq.com/articles/southern-q2-earnings-call-highlights ；https://quartr.com/events/the-southern-company-so-q2-2026_3YDwxjVN"),
       ("WELL FFO 指引 $6.36–6.44、股息 +15%", "SEC 8-K https://www.sec.gov/Archives/edgar/data/0000766704/000076670426000026/a2q26earningsrelease991.htm"),
       ("PLD Core FFO 指引 $6.22–6.30", "https://ir.prologis.com/news-events/press-releases/detail/1044/prologis-reports-second-quarter-2026-results"),
       ("EQIX AFFO 指引 $42.69–43.29", "https://investor.equinix.com/news-events/press-releases/detail/1114/equinix-reports-second-quarter-results-raises-2026"),
       ("LIN 2026 EPS 指引 $17.70–17.90、backlog $110 億", "SEC 8-K https://www.sec.gov/Archives/edgar/data/0001707925/000165495426007052/lin_ex991.htm"),
       ("NEM Q2 2026 FCF $22 億、AISC、回購、淨現金", "SEC 8-K https://www.sec.gov/Archives/edgar/data/0001164727/000116472726000034/newmontq22026earningsrelea.htm"),
       ("FCX Q3 2026 營運更新、Grasberg 恢復時間表", "SEC 8-K https://www.sec.gov/Archives/edgar/data/0000831259/000083125926000043/fcx_261002.htm"),
       ("關鍵假設", "rf=10Y 4.6%；ERP 4.5%；Beta、g1、gT、轉換率、目標/退出倍數均為分析師估計（見 Valuation/Picks 黃底格）；REIT 目標 P/FFO 為估計值。"),
       ("免責聲明", "本工作簿僅供研究與教育用途，不構成投資建議。部分數據來自二手媒體，可能存在誤差；投資決策前請核對公司原始披露。")]
for i, (a, b) in enumerate(SRC):
    r = 5 + i
    put(So, r, 1, i + 1, al=CENTER); put(So, r, 2, a, al=WRAP); put(So, r, 3, b, NOTE, al=WRAP)
    So.row_dimensions[r].height = 30


from openpyxl.chart import BarChart, Reference
ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"
ch.title = "板塊 Fwd P/E vs 10Y 均值"; ch.y_axis.title = "P/E (x)"
data = Reference(S, min_col=3, max_col=3, min_row=4, max_row=16); ch.add_data(data, titles_from_data=True)
data2 = Reference(S, min_col=5, max_col=5, min_row=4, max_row=16); ch.add_data(data2, titles_from_data=True)
ch.set_categories(Reference(S, min_col=2, min_row=5, max_row=16))
ch.height = 8; ch.width = 22
S.add_chart(ch, "U4")
ch2 = BarChart(); ch2.type = "bar"; ch2.title = "5 年預期年化回報"; ch2.y_axis.number_format = "0%"
ch2.add_data(Reference(P, min_col=14, max_col=14, min_row=5, max_row=17), titles_from_data=True)
ch2.set_categories(Reference(P, min_col=1, min_row=6, max_row=17))
ch2.height = 9; ch2.width = 16; ch2.legend = None
P.add_chart(ch2, "B26")

# order sheets
order = ["Cover", "Summary", "Sector_Map", "Methodology", "Competition", "Macro", "Comps", "Financials", "Valuation", "Scorecard", "Picks", "Sources"]
wb._sheets = [wb[n] for n in order]
for ws in wb.worksheets:
    ws.sheet_properties.tabColor = NAVY if ws.title in ("Cover", "Summary", "Picks") else MID
    if ws.title != "Cover":
        ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
wb.save(OUT)
print("saved", OUT)
